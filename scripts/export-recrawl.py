# -*- coding: utf-8 -*-
"""Список адресов на переобход — Excel в папку exports/.

Собирает страницы, изменённые в последних работах, тремя группами:
новые страницы подбора дробемётов, обновлённые хабы и страницы
ЛВМ/ЛГМ/ПГС/ХТС после редактуры. Адреса берутся из собранного
astro/dist — в список попадает только то, что реально отдаёт сервер.

Запуск (из корня проекта, после npm run build):
    python scripts/export-recrawl.py
"""
import datetime
import io
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SITE = 'https://uralinduktor.ru'
DIST = 'astro/dist'
OUT_DIR = 'exports'

GROUPS = [
    ('Подбор дробемётов — новые страницы', [
        '/products/drobemetnoe-oborudovanie/podbor-drobemeta/',
        '/products/drobemetnoe-oborudovanie/dlya-ochistki-otlivok/',
        '/products/drobemetnoe-oborudovanie/dlya-metallokonstruktsiy/',
        '/products/drobemetnoe-oborudovanie/dlya-pokovok/',
    ]),
    ('Обновлённые хабы — метатеги, таблицы, FAQ', [
        '/products/induktsionnye-plavilnye-pechi/',
        '/products/drobemetnoe-oborudovanie/',
        '/products/induktsionnye-nagrevateli/vysokochastotnye/',
    ]),
    # ЛВМ / ЛГМ / ПГС / ХТС — весь раздел после редактуры текстов
    ('ЛВМ, ЛГМ, ПГС, ХТС — переписанные тексты', None),
]

SECTIONS = ['products/oborudovanie-lvm', 'products/oborudovanie-lgm',
            'products/oborudovanie-dlya-peschano-glinistyh-smesej-pgs',
            'products/linii-dlya-holodnotverdeyushhih-smesej-hts']


def section_pages():
    urls = []
    for sec in SECTIONS:
        for root, _, files in os.walk(os.path.join(DIST, sec)):
            if 'index.html' not in files:
                continue
            urls.append('/' + os.path.relpath(root, DIST).replace(os.sep, '/') + '/')
    return sorted(urls)


def page_title(url):
    p = os.path.join(DIST, url.strip('/'), 'index.html')
    if not os.path.isfile(p):
        return None
    m = re.search(r'<title>(.*?)</title>', io.open(p, encoding='utf-8').read(), re.S)
    return (m.group(1).strip() if m else '')


def main():
    if not os.path.isdir(DIST):
        sys.exit('нет %s — сначала cd astro && npm run build' % DIST)

    rows, missing = [], []
    for name, urls in GROUPS:
        for u in (urls if urls is not None else section_pages()):
            t = page_title(u)
            if t is None:
                missing.append(u)
                continue
            rows.append((SITE + u, name, t))

    if missing:
        print('ВНИМАНИЕ, этих страниц нет в сборке:')
        for u in missing:
            print('   ' + u)

    wb = Workbook()
    head_font = Font(bold=True, color='FFFFFF')
    head_fill = PatternFill('solid', fgColor='2F6B4F')

    ws = wb.active
    ws.title = 'На переобход'
    ws.append(['URL', 'Группа', 'Заголовок страницы'])
    for r in rows:
        ws.append(list(r))
    for c in range(1, 4):
        cell = ws.cell(row=1, column=c)
        cell.font, cell.fill = head_font, head_fill
        cell.alignment = Alignment(vertical='center')
    for c, w in ((1, 95), (2, 42), (3, 80)):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = 'A1:C%d' % (len(rows) + 1)

    sw = wb.create_sheet('Сводка')
    sw.append(['Группа', 'Страниц'])
    for name, _ in GROUPS:
        sw.append([name, sum(1 for r in rows if r[1] == name)])
    sw.append(['ИТОГО', len(rows)])
    for c in range(1, 3):
        sw.cell(row=1, column=c).font = head_font
        sw.cell(row=1, column=c).fill = head_fill
    sw.cell(row=sw.max_row, column=1).font = Font(bold=True)
    sw.cell(row=sw.max_row, column=2).font = Font(bold=True)
    sw.column_dimensions['A'].width = 46
    sw.column_dimensions['B'].width = 12

    os.makedirs(OUT_DIR, exist_ok=True)
    stamp = datetime.date.today().isoformat()
    xlsx = os.path.join(OUT_DIR, 'recrawl-%s.xlsx' % stamp)
    wb.save(xlsx)
    txt = os.path.join(OUT_DIR, 'recrawl-%s.txt' % stamp)
    io.open(txt, 'w', encoding='utf-8', newline='\n').write('\n'.join(r[0] for r in rows) + '\n')

    for name, _ in GROUPS:
        print('%-44s %3d' % (name, sum(1 for r in rows if r[1] == name)))
    print('ИТОГО: %d' % len(rows))
    print('  %s' % xlsx)
    print('  %s (тот же список простым текстом)' % txt)


if __name__ == '__main__':
    main()
